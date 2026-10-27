"""
Complete Batch Pipeline to import the 2 WhatsApp ZIP files into:
  1. dataset/Wrist_Dataset/{id}/ (4 standard image files + metadata.json)
  2. Supabase Storage and database tables
  3. doctor_validation_prakriti_master.xlsx (with embedded photo thumbnails)
"""

import sys
import os
import json
import zipfile
import shutil
import io
from pathlib import Path
from PIL import Image as PILImage
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage

PROJECT_ROOT = Path("c:/Users/hp/Desktop/ayurveda prakriti classifier")
WRIST_DIR = PROJECT_ROOT / "dataset" / "Wrist_Dataset"
MASTER_XLSX = PROJECT_ROOT / "doctor_validation_prakriti_master.xlsx"
DATASET_MASTER_XLSX = PROJECT_ROOT / "dataset" / "doctor_validation_prakriti_master.xlsx"

sys.path.insert(0, str(PROJECT_ROOT))
sys.path.insert(0, str(PROJECT_ROOT / "mini"))
sys.path.insert(0, str(PROJECT_ROOT / "whatsapp_pipeline"))

from parser import parse_chat_log, group_by_subject
from extractor import extract_details
from view_classifier import classify_hand_views
from predictor import PrakritiPredictor
from supabasehelper import upload_complete_analysis, supabase_client

VIEW_TO_FILENAME = {
    "palm_open_front": "palm_open_front.jpg",
    "palm_open_back": "palm_open_back.jpg",
    "palm_closed_front": "palm_close_front.jpg",
    "palm_closed_back": "palm_close_back.jpg"
}

def get_next_id():
    max_num = 1000
    for d in WRIST_DIR.iterdir():
        if d.is_dir():
            num_str, _, _ = d.name.partition("_")
            try:
                max_num = max(max_num, int(num_str))
            except ValueError:
                pass
    return max_num + 1

def add_thumbnail_to_cell(ws, img_path, cell_coord):
    try:
        img = PILImage.open(img_path)
        img.thumbnail((90, 80))
        buf = io.BytesIO()
        img.convert("RGB").save(buf, format="JPEG", quality=85)
        buf.seek(0)
        ox_img = OpenpyxlImage(buf)
        ox_img.width = img.width
        ox_img.height = img.height
        ws.add_image(ox_img, cell_coord)
        return True
    except Exception as e:
        print(f"  [Thumbnail Warning] Could not add image {img_path.name}: {e}")
        return False

def run_import():
    zip_files = [
        PROJECT_ROOT / "zip" / "WhatsApp Chat - Palm dataset 2.zip",
        PROJECT_ROOT / "zip" / "WhatsApp Chat - Palm dataset 3.zip"
    ]
    
    predictor = PrakritiPredictor()
    current_id = get_next_id()
    print(f"Starting numeric participant ID: {current_id}")
    
    # Staging
    staging_base = PROJECT_ROOT / "staging_zip_import"
    staging_base.mkdir(parents=True, exist_ok=True)
    
    all_processed_subjects = []
    
    for zpath in zip_files:
        if not zpath.exists():
            print(f"[Error] Missing ZIP: {zpath}")
            continue
            
        print(f"\n=======================================================")
        print(f"Extracting & Parsing Archive: {zpath.name}")
        print(f"=======================================================")
        
        extract_dir = staging_base / zpath.stem
        extract_dir.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(zpath, 'r') as zf:
            zf.extractall(extract_dir)
            
        chat_file = extract_dir / "_chat.txt"
        if not chat_file.exists():
            print(f"[Error] _chat.txt missing in {extract_dir}")
            continue
            
        messages = parse_chat_log(chat_file)
        subjects = group_by_subject(messages)
        print(f"Found {len(subjects)} subjects in {zpath.name}.")
        
        for s_idx, subj in enumerate(subjects, 1):
            text = subj["associated_text"]
            images = subj["images"]
            details = extract_details(text)
            
            # Resolve image paths
            full_img_paths = []
            for img_name in images:
                p = extract_dir / img_name
                if p.exists():
                    full_img_paths.append(str(p))
                else:
                    # Case-insensitive
                    for c in extract_dir.iterdir():
                        if c.name.lower() == img_name.lower():
                            full_img_paths.append(str(c))
                            break
                            
            if len(full_img_paths) != 4:
                print(f"  [Warning] Subject {s_idx} ({details.get('name')}) has {len(full_img_paths)} images instead of 4. Skipping.")
                continue
                
            # Classify views
            print(f"\nProcessing [{s_idx}/{len(subjects)} in {zpath.name}]: '{details.get('name')}' (Age: {details.get('age')}, Gen: {details.get('gender')})")
            view_mapping = classify_hand_views(full_img_paths)
            mapped_views = {cat: path for path, cat in view_mapping.items() if cat}
            
            gender_code = "M" if str(details.get("gender", "")).lower().startswith("m") else ("F" if str(details.get("gender", "")).lower().startswith("f") else "Unknown")
            gender_str = "male" if gender_code == "M" else ("female" if gender_code == "F" else "unknown")
            subj_id_name = f"{current_id}_{gender_str}"
            
            # Destination folder in Wrist_Dataset
            subj_dir = WRIST_DIR / subj_id_name
            subj_dir.mkdir(parents=True, exist_ok=True)
            
            saved_paths = {}
            for cat, fname in VIEW_TO_FILENAME.items():
                src_img = mapped_views.get(cat)
                if src_img and Path(src_img).exists():
                    dst_img = subj_dir / fname
                    shutil.copy2(src_img, dst_img)
                    saved_paths[cat] = dst_img
                    
            # Write metadata.json
            meta = {
                "id": f"person_{current_id}",
                "name": details.get("name"),
                "gender": gender_str,
                "age": details.get("age"),
                "vein_visibility": details.get("vein_visibility"),
                "mole": details.get("mole"),
                "skin_texture": details.get("skin_texture"),
                "source_zip": zpath.name
            }
            (subj_dir / "metadata.json").write_text(json.dumps(meta, indent=2))
            
            # Run ML Predictor for reference
            pred_res = predictor.predict_subject(
                age=details.get("age"),
                gender=gender_code,
                palm_open_front_path=str(saved_paths.get("palm_open_front", "")),
                palm_open_back_path=str(saved_paths.get("palm_open_back", "")),
                palm_closed_front_path=str(saved_paths.get("palm_closed_front", "")),
                palm_closed_back_path=str(saved_paths.get("palm_closed_back", "")),
                details=details
            )
            
            ml_dosha = pred_res.get("predicted_dosha", "Unknown")
            ml_pct = f"k-{pred_res.get('kapha_prob', 0)*100:.0f}%, p-{pred_res.get('pitta_prob', 0)*100:.0f}%, v-{pred_res.get('vata_prob', 0)*100:.0f}%"
            
            # Upload to Supabase
            if supabase_client:
                try:
                    pil_imgs = {}
                    for cat, p in saved_paths.items():
                        pil_imgs[cat] = PILImage.open(p)
                    
                    probs_dict = {
                        "Kapha": pred_res.get("kapha_prob", 0.0),
                        "Pitta": pred_res.get("pitta_prob", 0.0),
                        "Vata": pred_res.get("vata_prob", 0.0)
                    }
                    sb_res = upload_complete_analysis(
                        user_id=subj_id_name,
                        gender=gender_str,
                        images_dict=pil_imgs,
                        probabilities=probs_dict,
                        final_dosha=ml_dosha,
                        confidence=pred_res.get("confidence", 1.0),
                        age=int(details["age"]) if details.get("age") not in (None, "Unknown") else None,
                        skin_texture=details.get("skin_texture", "Smooth"),
                        vein_visibility=details.get("vein_visibility", "Low"),
                        mole=1 if details.get("mole") == "Yes" else 0,
                        name=details.get("name")
                    )
                    if sb_res.get("success"):
                        print(f"  [Supabase] Uploaded '{details.get('name')}' as ID '{subj_id_name}'")
                    else:
                        print(f"  [Supabase] Notice: {sb_res.get('errors')}")
                except Exception as sb_err:
                    print(f"  [Supabase Error]: {sb_err}")
                    
            record_info = {
                "id": subj_id_name,
                "name": details.get("name") or "Unknown",
                "age": details.get("age"),
                "sex": gender_code,
                "vein": details.get("vein_visibility", "Low"),
                "mole": details.get("mole", "No"),
                "skin": details.get("skin_texture", "Smooth"),
                "ml_dosha": ml_dosha,
                "ml_pct": ml_pct,
                "saved_paths": saved_paths
            }
            all_processed_subjects.append(record_info)
            current_id += 1
            
    print(f"\nTotal new subjects extracted and organized: {len(all_processed_subjects)}")
    
    # Append to doctor_validation_prakriti_master.xlsx
    append_to_excel(all_processed_subjects)

def append_to_excel(subjects):
    print("\nAppending new participants to doctor_validation_prakriti_master.xlsx...")
    wb = openpyxl.load_workbook(MASTER_XLSX)
    ws = wb["Doctor Validation Master"]
    
    # Get column mapping from row 1
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    col_map = {name: c for c, name in enumerate(headers, 1) if name}
    
    fill_pending = PatternFill(start_color="FFF9E6", end_color="FFF9E6", fill_type="solid")
    f_data = Font(name="Calibri", size=10)
    f_bold = Font(name="Calibri", size=10, bold=True)
    border_thin = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    
    start_row = ws.max_row + 1
    
    for idx, s in enumerate(subjects):
        r_num = start_row + idx
        ws.row_dimensions[r_num].height = 80
        
        vv = s["vein"]
        mole = s["mole"]
        st = s["skin"]
        
        # Binary one-hot mapping for the 3 available features
        vv_low = 1 if vv == "Low" else 0
        vv_med = 1 if vv == "Med" else 0
        vv_high = 1 if vv == "High" else 0
        
        mole_val = 1 if mole == "Yes" else 0
        
        st_dry = 1 if st == "Dry" else 0
        st_rough = 1 if st == "Rough" else 0
        st_smooth = 1 if st == "Smooth" else 0
        st_oily = 1 if st == "Oily" else 0
        
        row_vals = {
            "Participant ID": s["id"],
            "Participant Name": s["name"],
            "Age": s["age"] if s["age"] != "Unknown" else None,
            "Sex": s["sex"],
            "Validation Status": "Pending Doctor Validation",
            "Doctor Verified Prakriti (Ground Truth)": None,
            "Doctor Vein Visibility": vv,
            "Doctor Mole Presence": mole,
            "Doctor Skin Texture": st,
            "Doctor Skin Color": None,
            "Doctor Palm Color": None,
            "Doctor Hair Thickness": None,
            "Doctor Hair Color": None,
            "Doctor Clinical Notes": "New cohort from WhatsApp",
            "ML Predicted Dosha": s["ml_dosha"],
            "ML Confidence / Percentages": s["ml_pct"],
            "SC_Fair": 0, "SC_Medium": 0, "SC_Dark": 0,
            "ST_Rough": st_rough, "ST_Smooth": st_smooth, "ST_Dry": st_dry, "ST_Oily": st_oily,
            "SH_Thick": 0, "SH_Thin": 0, "SH_Med": 0,
            "SHC_Brown": 0, "SHC_Black": 0,
            "Mole on skin": mole_val,
            "VV_Low": vv_low, "VV_Med": vv_med, "VV_High": vv_high,
            "PC_PinkRed": 0, "PC_PinkWhite": 0, "PC_Brown": 0
        }
        
        for col_name, val in row_vals.items():
            if col_name in col_map:
                c_idx = col_map[col_name]
                cell = ws.cell(row=r_num, column=c_idx, value=val)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border_thin
                cell.font = f_bold if col_name in ["Participant ID", "Validation Status"] else f_data
                
                # Highlight doctor columns
                if col_name.startswith("Doctor") or col_name == "Validation Status":
                    cell.fill = fill_pending
                    
        # Add embedded photo thumbnails
        VIEW_COLS = {
            "palm_open_front": "E",
            "palm_open_back": "F",
            "palm_closed_front": "G",
            "palm_closed_back": "H"
        }
        for view, col_letter in VIEW_COLS.items():
            img_p = s["saved_paths"].get(view)
            if img_p and Path(img_p).exists():
                add_thumbnail_to_cell(ws, Path(img_p), f"{col_letter}{r_num}")
                
    wb.save(MASTER_XLSX)
    print(f"Successfully saved updated master tracking sheet -> {MASTER_XLSX}")
    
    # Mirror copy in dataset/
    shutil.copy2(MASTER_XLSX, DATASET_MASTER_XLSX)
    print(f"Mirrored copy saved -> {DATASET_MASTER_XLSX}")
    print(f"Total rows now in master tracking sheet: {ws.max_row}")

if __name__ == "__main__":
    run_import()
