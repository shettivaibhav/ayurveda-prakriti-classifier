"""
Sync Unified Records Script
Synchronizes doctor-verified clinical data across:
  1. Master Excel sheets (dataset/doctor_validation_prakriti_master.xlsx and root)
     - Sets Validation Status = 'Verified (Clinical Ground Truth)' for rows 95-121 (except row 102)
     - Recomputes all 19 one-hot binary clinical feature columns
     - Handles Row 102 (1101_Female) as a duplicate repeat of Row 101 (1100_Female)
  2. Local participant folders (dataset/Wrist_Dataset/{id}/metadata.json)
  3. Supabase Cloud Database (wrist_uploads table, storing complete ground truth entries with full feature JSON)
"""

import sys
import json
import shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

PROJECT_ROOT = Path("c:/Users/hp/Desktop/ayurveda prakriti classifier")
SRC_XLSX = PROJECT_ROOT / "dataset" / "doctor_validation_prakriti_master.xlsx"
ROOT_XLSX = PROJECT_ROOT / "doctor_validation_prakriti_master.xlsx"
WRIST_DIR = PROJECT_ROOT / "dataset" / "Wrist_Dataset"

sys.path.insert(0, str(PROJECT_ROOT / "mini"))
from supabasehelper import supabase_client


def sync_excel_and_local():
    print(f"Loading master sheet: {SRC_XLSX}...")
    wb = openpyxl.load_workbook(SRC_XLSX)
    ws = wb.active

    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    col_map = {name: c for c, name in enumerate(headers, 1) if name}

    fill_verified = PatternFill(start_color="F2F8F2", end_color="F2F8F2", fill_type="solid")  # Soft green
    fill_duplicate = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid") # Soft gray
    f_bold = Font(name="Calibri", size=10, bold=True)
    f_italic = Font(name="Calibri", size=10, italic=True)

    verified_records = []
    duplicate_records = []

    print("\n--- Processing rows 95 to 121 in Excel ---")
    for r in range(95, 122):
        row_id = str(ws.cell(r, col_map["Participant ID"]).value).strip()
        name = str(ws.cell(r, col_map["Participant Name"]).value).strip()

        # Check for row 102 (Repeat of 101)
        if r == 102 or row_id.lower().startswith("1101"):
            print(f"  [ROW {r:3d}] Flagging {row_id} ({name}) as DUPLICATE (Repeat of 1100_Female)...")
            ws.cell(r, col_map["Validation Status"], value="Duplicate (Repeat of 1100_Female)")
            ws.cell(r, col_map["Doctor Verified Prakriti (Ground Truth)"], value="V-K (Duplicate of 1100)")
            ws.cell(r, col_map["Doctor Clinical Notes"], value="Duplicate record: Repeat of 1100_Female (Janita David). Excluded from training set.")

            # Style row
            for c in range(1, ws.max_column + 1):
                ws.cell(r, c).fill = fill_duplicate
            ws.cell(r, col_map["Validation Status"]).font = f_italic

            dup_info = {
                "row": r,
                "id": row_id,
                "name": name,
                "validation_status": "Duplicate (Repeat of 1100_Female)",
                "is_duplicate": True,
                "duplicate_of": "1100_female",
                "notes": "Repeat of 1100_Female (Janita David)"
            }
            duplicate_records.append(dup_info)
            continue

        # Standard verified record
        prakriti = str(ws.cell(r, col_map["Doctor Verified Prakriti (Ground Truth)"]).value).strip()
        vv = str(ws.cell(r, col_map["Doctor Vein Visibility"]).value or "").strip()
        mole = str(ws.cell(r, col_map["Doctor Mole Presence"]).value or "").strip()
        st = str(ws.cell(r, col_map["Doctor Skin Texture"]).value or "").strip()
        sc = str(ws.cell(r, col_map["Doctor Skin Color"]).value or "").strip()
        pc = str(ws.cell(r, col_map["Doctor Palm Color"]).value or "").strip()
        sh = str(ws.cell(r, col_map["Doctor Hair Thickness"]).value or "").strip()
        shc = str(ws.cell(r, col_map["Doctor Hair Color"]).value or "").strip()

        # 1. Update Validation Status and Notes
        ws.cell(r, col_map["Validation Status"], value="Verified (Clinical Ground Truth)")
        ws.cell(r, col_map["Doctor Clinical Notes"], value="Clinically verified by Doctor in Cohort 2")

        # 2. Recompute 19 One-Hot Binary Columns
        # Skin Color (SC)
        ws.cell(r, col_map["SC_Fair"], value=1 if sc == "Fair" else 0)
        ws.cell(r, col_map["SC_Medium"], value=1 if sc == "Medium" else 0)
        ws.cell(r, col_map["SC_Dark"], value=1 if sc == "Dark" else 0)

        # Skin Texture (ST)
        ws.cell(r, col_map["ST_Rough"], value=1 if st == "Rough" else 0)
        ws.cell(r, col_map["ST_Smooth"], value=1 if st == "Smooth" else 0)
        ws.cell(r, col_map["ST_Dry"], value=1 if st == "Dry" else 0)
        ws.cell(r, col_map["ST_Oily"], value=1 if st == "Oily" else 0)

        # Skin Hair Thickness (SH)
        ws.cell(r, col_map["SH_Thick"], value=1 if sh == "Thick" else 0)
        ws.cell(r, col_map["SH_Thin"], value=1 if sh == "Thin" else 0)
        ws.cell(r, col_map["SH_Med"], value=1 if sh == "Med" else 0)

        # Skin Hair Color (SHC)
        ws.cell(r, col_map["SHC_Brown"], value=1 if shc == "Brown" else 0)
        ws.cell(r, col_map["SHC_Black"], value=1 if shc == "Black" else 0)

        # Mole
        ws.cell(r, col_map["Mole on skin"], value=1 if mole == "Yes" else 0)

        # Vein Visibility (VV)
        ws.cell(r, col_map["VV_Low"], value=1 if vv == "Low" else 0)
        ws.cell(r, col_map["VV_Med"], value=1 if vv == "Med" else 0)
        ws.cell(r, col_map["VV_High"], value=1 if vv == "High" else 0)

        # Palm Color (PC)
        ws.cell(r, col_map["PC_PinkRed"], value=1 if pc == "PinkRed" else 0)
        ws.cell(r, col_map["PC_PinkWhite"], value=1 if pc == "PinkWhite" else 0)
        ws.cell(r, col_map["PC_Brown"], value=1 if pc == "Brown" else 0)

        # Apply soft green styling to indicate clinical verification
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).fill = fill_verified
        ws.cell(r, col_map["Validation Status"]).font = f_bold
        ws.cell(r, col_map["Doctor Verified Prakriti (Ground Truth)"]).font = f_bold

        rec_info = {
            "row": r,
            "id": row_id,
            "name": name,
            "doctor_verified_prakriti": prakriti,
            "validation_status": "Verified (Clinical Ground Truth)",
            "features": {
                "vein_visibility": vv,
                "mole": mole,
                "skin_texture": st,
                "skin_color": sc,
                "palm_color": pc,
                "hair_thickness": sh,
                "hair_color": shc,
            },
            "is_duplicate": False
        }
        verified_records.append(rec_info)

    # Save to both dataset/ and root
    wb.save(SRC_XLSX)
    shutil.copy2(SRC_XLSX, ROOT_XLSX)
    print(f"Excel saved to {SRC_XLSX} and synced to {ROOT_XLSX}")

    # 3. Update local metadata.json in dataset/Wrist_Dataset/
    for rec in verified_records:
        folder_name = rec["id"].lower()
        target_dir = WRIST_DIR / folder_name
        if not target_dir.exists():
            for d in WRIST_DIR.iterdir():
                if d.is_dir() and d.name.lower() == folder_name:
                    target_dir = d
                    break
        if target_dir.exists():
            meta_path = target_dir / "metadata.json"
            meta = {}
            if meta_path.exists():
                try:
                    meta = json.loads(meta_path.read_text())
                except Exception:
                    pass
            meta["doctor_verified_prakriti"] = rec["doctor_verified_prakriti"]
            meta["validation_status"] = rec["validation_status"]
            meta["doctor_verified_features"] = rec["features"]
            meta["doctor_notes"] = "Clinically verified by Doctor in Cohort 2"
            meta["is_duplicate"] = False
            meta_path.write_text(json.dumps(meta, indent=4))

    # Update duplicate record metadata
    for dup in duplicate_records:
        folder_name = dup["id"].lower()
        target_dir = WRIST_DIR / folder_name
        if target_dir.exists():
            meta_path = target_dir / "metadata.json"
            meta = {}
            if meta_path.exists():
                try:
                    meta = json.loads(meta_path.read_text())
                except Exception:
                    pass
            meta["validation_status"] = "Duplicate (Repeat of 1100_Female)"
            meta["is_duplicate"] = True
            meta["duplicate_of"] = "1100_female"
            meta["doctor_notes"] = "Duplicate record - repeat of 1100_Female (Janita David). Excluded from training set."
            meta_path.write_text(json.dumps(meta, indent=4))

    return verified_records, duplicate_records


def sync_to_supabase(verified_records, duplicate_records):
    print("\n--- Syncing records to Supabase Cloud Database ---")
    if not supabase_client:
        print("  [Supabase Warning] supabase_client not initialized.")
        return

    # Check which user_ids already have a ground truth entry in Supabase
    try:
        existing = supabase_client.table("wrist_uploads").select("user_id").eq("angle", "doctor_ground_truth_cohort2").execute()
        existing_uids = {r["user_id"] for r in (existing.data or [])}
    except Exception as e:
        existing_uids = set()

    synced_count = 0

    # 1. Sync verified records
    for rec in verified_records:
        uid = rec["id"]
        if uid in existing_uids:
            continue

        gender = "female" if "female" in uid.lower() else "male"
        payload = {
            "user_id": uid,
            "gender": gender,
            "angle": "doctor_ground_truth_cohort2",
            "predicted_dosha": rec["doctor_verified_prakriti"],
            "confidence": 1.0,
            "image_data": json.dumps({
                "validation_status": rec["validation_status"],
                "doctor_verified_prakriti": rec["doctor_verified_prakriti"],
                "clinical_features": rec["features"],
                "doctor_notes": "Clinically verified by Doctor in Cohort 2",
                "is_duplicate": False
            })
        }
        try:
            res = supabase_client.table("wrist_uploads").insert(payload).execute()
            if res.data:
                synced_count += 1
                print(f"  [Supabase] Recorded {uid} (Dosha={rec['doctor_verified_prakriti']})")
        except Exception as e:
            print(f"  [Supabase Error for {uid}]: {e}")

    # 2. Sync duplicate record
    for dup in duplicate_records:
        uid = dup["id"]
        if uid in existing_uids:
            continue
        gender = "female" if "female" in uid.lower() else "male"
        payload = {
            "user_id": uid,
            "gender": gender,
            "angle": "doctor_ground_truth_cohort2",
            "predicted_dosha": "Duplicate",
            "confidence": 0.0,
            "image_data": json.dumps({
                "validation_status": dup["validation_status"],
                "is_duplicate": True,
                "duplicate_of": dup["duplicate_of"],
                "doctor_notes": dup["notes"]
            })
        }
        try:
            res = supabase_client.table("wrist_uploads").insert(payload).execute()
            if res.data:
                synced_count += 1
                print(f"  [Supabase] Recorded {uid} as DUPLICATE.")
        except Exception as e:
            print(f"  [Supabase Error for duplicate {uid}]: {e}")

    print(f"Successfully synced {synced_count} doctor-verified records to Supabase!")


if __name__ == "__main__":
    verified, duplicate = sync_excel_and_local()
    sync_to_supabase(verified, duplicate)
    print("\n[OK] Sync completed across Excel, Local Folders, and Supabase!")
