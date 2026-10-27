"""
Supabase helper functions for uploading wrist images and data
"""
import io
from PIL import Image
from supabase import create_client
from datetime import datetime
import hashlib
import streamlit as st


# Supabase configuration
SUPABASE_URL = None
SUPABASE_KEY = None
supabase_client = None

def _load_secrets():
    global SUPABASE_URL, SUPABASE_KEY, supabase_client
    # 1. Try Streamlit secrets
    try:
        SUPABASE_URL = st.secrets["supabase"]["url"]
        SUPABASE_KEY = st.secrets["supabase"]["key"]
    except Exception:
        pass
        
    # 2. Try parsing secrets.toml file directly
    if not SUPABASE_URL or not SUPABASE_KEY:
        try:
            import tomllib
            from pathlib import Path
            current_dir = Path(__file__).resolve().parent
            toml_paths = [
                current_dir / ".streamlit" / "secrets.toml",
                current_dir.parent / ".streamlit" / "secrets.toml",
                Path.cwd() / ".streamlit" / "secrets.toml",
                Path.cwd() / "mini" / ".streamlit" / "secrets.toml"
            ]
            for path in toml_paths:
                if path.exists():
                    with open(path, "rb") as f:
                        data = tomllib.load(f)
                        if "supabase" in data:
                            SUPABASE_URL = data["supabase"].get("url")
                            SUPABASE_KEY = data["supabase"].get("key")
                            break
        except Exception:
            pass

    if SUPABASE_URL and SUPABASE_KEY:
        try:
            supabase_client = create_client(SUPABASE_URL, SUPABASE_KEY)
        except Exception as e:
            print(f"[Supabase Helper] Failed to initialize client: {e}")

_load_secrets()

def get_gemini_api_key():
    """Get Gemini API Key from Streamlit secrets or secrets.toml file directly"""
    # 1. Try from Streamlit secrets
    try:
        if "supabase" in st.secrets and "gemini_api_key" in st.secrets["supabase"]:
            return st.secrets["supabase"]["gemini_api_key"]
        if "gemini_api_key" in st.secrets:
            return st.secrets["gemini_api_key"]
    except Exception:
        pass

    # 2. Try parsing secrets.toml file directly
    try:
        import tomllib
        from pathlib import Path
        current_dir = Path(__file__).resolve().parent
        toml_paths = [
            current_dir / ".streamlit" / "secrets.toml",
            current_dir.parent / ".streamlit" / "secrets.toml",
            Path.cwd() / ".streamlit" / "secrets.toml",
            Path.cwd() / "mini" / ".streamlit" / "secrets.toml"
        ]
        for path in toml_paths:
            if path.exists():
                with open(path, "rb") as f:
                    data = tomllib.load(f)
                    if "supabase" in data and "gemini_api_key" in data["supabase"]:
                        return data["supabase"]["gemini_api_key"]
                    if "gemini_api_key" in data:
                        return data["gemini_api_key"]
    except Exception:
        pass
    return None


def get_groq_api_key():
    """Get Groq API Key from Streamlit secrets or secrets.toml file directly"""
    # 1. Try from Streamlit secrets
    try:
        if "supabase" in st.secrets:
            for k, v in st.secrets["supabase"].items():
                if k.lower() == "groq_api_key":
                    return v
        for k, v in st.secrets.items():
            if k.lower() == "groq_api_key":
                return v
    except Exception:
        pass

    # 2. Try parsing secrets.toml file directly
    try:
        import tomllib
        from pathlib import Path
        current_dir = Path(__file__).resolve().parent
        toml_paths = [
            current_dir / ".streamlit" / "secrets.toml",
            current_dir.parent / ".streamlit" / "secrets.toml",
            Path.cwd() / ".streamlit" / "secrets.toml",
            Path.cwd() / "mini" / ".streamlit" / "secrets.toml"
        ]
        for path in toml_paths:
            if path.exists():
                with open(path, "rb") as f:
                    data = tomllib.load(f)
                    # Check inside [supabase] table case-insensitively
                    if "supabase" in data:
                        for k, v in data["supabase"].items():
                            if k.lower() == "groq_api_key":
                                return v
                    # Check at root level case-insensitively
                    for k, v in data.items():
                        if k.lower() == "groq_api_key":
                            return v
    except Exception:
        pass
    return None


def compress_image(image, max_size_kb=200, quality=85):
    """
    Compress PIL Image to target size
    
    Args:
        image: PIL Image object
        max_size_kb: Maximum size in KB (default 200KB)
        quality: Initial JPEG quality (default 85)
    
    Returns:
        bytes: Compressed image bytes
    """
    # Convert to RGB if necessary
    if image.mode in ('RGBA', 'P'):
        image = image.convert('RGB')
    
    # Start with initial quality
    output = io.BytesIO()
    image.save(output, format='JPEG', quality=quality, optimize=True)
    
    # Reduce quality if still too large
    while output.tell() > max_size_kb * 1024 and quality > 20:
        output = io.BytesIO()
        quality -= 5
        image.save(output, format='JPEG', quality=quality, optimize=True)
    
    return output.getvalue()


def upload_image_to_storage(image, user_id, gender, angle_key):
    """
    Upload a single image to Supabase storage
    
    Args:
        image: PIL Image object
        user_id: User ID (e.g., "1001")
        gender: Gender ("male" or "female")
        angle_key: Angle key (e.g., "palm_open_front")
    
    Returns:
        dict: {success: bool, url: str or None, error: str or None}
    """
    try:
        # Compress image
        img_bytes = compress_image(image, max_size_kb=200, quality=85)
        
        # Create storage path: {id}_{gender}/{angle}.jpg
        folder_name = f"{user_id}_{gender}"
        filename = f"{angle_key}.jpg"
        storage_path = f"{folder_name}/{filename}"
        
        # Upload to storage bucket
        response = supabase_client.storage.from_("wrist_images").upload(
            storage_path,
            img_bytes,
            {"content-type": "image/jpeg", "upsert": "true"}
        )
        
        # Check for errors
        if hasattr(response, 'error') and response.error:
            return {
                "success": False,
                "url": None,
                "error": f"Storage error: {response.error.message}"
            }
        
        # Get public URL
        public_url = supabase_client.storage.from_("wrist_images").get_public_url(storage_path)
        
        return {
            "success": True,
            "url": public_url,
            "error": None,
            "size_kb": len(img_bytes) / 1024
        }
        
    except Exception as e:
        return {
            "success": False,
            "url": None,
            "error": str(e)
        }


def upload_all_images(images_dict, user_id, gender):
    """
    Upload all 4 wrist images to Supabase storage
    
    Args:
        images_dict: Dict with keys like "palm_open_front" and PIL Image values
        user_id: User ID
        gender: Gender
    
    Returns:
        dict: {success: bool, urls: dict, errors: list}
    """
    results = {
        "success": True,
        "urls": {},
        "errors": []
    }
    
    for angle_key, image in images_dict.items():
        result = upload_image_to_storage(image, user_id, gender, angle_key)
        
        if result["success"]:
            results["urls"][angle_key] = result["url"]
        else:
            results["success"] = False
            results["errors"].append(f"{angle_key}: {result['error']}")
    
    return results


def insert_wrist_data(user_id, gender, probabilities, final_dosha, confidence, metadata_dict=None):
    """
    Insert wrist analysis data into Supabase table
    
    Args:
        user_id: User ID
        gender: Gender
        probabilities: Dict with "Vata", "Pitta", "Kapha" probabilities
        final_dosha: Final predicted dosha string (e.g., "Vata", "Vata-Pitta")
        confidence: Confidence score (0-1)
        metadata_dict: Dict with full clinical features, dominant/subdominant doshas, and doctor mode
    
    Returns:
        dict: {success: bool, record_id: int or None, error: str or None}
    """
    try:
        import json
        image_data_str = json.dumps(metadata_dict) if metadata_dict else None
        
        # Prepare data for insertion
        data = {
            "user_id": str(user_id),
            "gender": gender,
            "angle": "combined_4_views",
            "vata_prob": round(probabilities.get("Vata", 0.0), 4),
            "pitta_prob": round(probabilities.get("Pitta", 0.0), 4),
            "kapha_prob": round(probabilities.get("Kapha", 0.0), 4),
            "predicted_dosha": final_dosha,
            "confidence": round(confidence, 4),
            "image_data": image_data_str,
            "image_size": None
        }
        
        # Insert into table
        response = supabase_client.table("wrist_uploads").insert(data).execute()
        
        # Check for successful insert
        if response.data and len(response.data) > 0:
            record_id = response.data[0].get("id")
            return {
                "success": True,
                "record_id": record_id,
                "error": None
            }
        else:
            return {
                "success": False,
                "record_id": None,
                "error": "No data returned from insert"
            }
            
    except Exception as e:
        return {
            "success": False,
            "record_id": None,
            "error": str(e)
        }


def update_wrist_data_prakriti(user_id, predicted_dosha, probabilities, confidence):
    """
    Update predicted dosha and probabilities for an existing user_id record in Supabase
    """
    try:
        data = {
            "predicted_dosha": predicted_dosha,
            "vata_prob": round(probabilities.get("Vata", 0.0), 4),
            "pitta_prob": round(probabilities.get("Pitta", 0.0), 4),
            "kapha_prob": round(probabilities.get("Kapha", 0.0), 4),
            "confidence": round(confidence, 4)
        }
        
        response = supabase_client.table("wrist_uploads").update(data).eq("user_id", str(user_id)).execute()
        
        if response.data and len(response.data) > 0:
            return {
                "success": True,
                "error": None
            }
        else:
            return {
                "success": False,
                "error": f"No records updated for user_id '{user_id}'"
            }
    except Exception as e:
        return {
            "success": False,
            "error": str(e)
        }


def save_to_local_dataset(user_id, gender, images_dict, metadata_dict):
    """
    Saves high-res photos and metadata.json to dataset/Wrist_Dataset/{user_id}_{gender}/
    """
    try:
        import json
        from pathlib import Path
        project_root = Path(__file__).resolve().parent.parent
        target_dir = project_root / "dataset" / "Wrist_Dataset" / f"{user_id}_{gender}"
        target_dir.mkdir(parents=True, exist_ok=True)
        
        view_filenames = {
            "palm_open_front": "palm_open_front.jpg",
            "palm_open_back": "palm_open_back.jpg",
            "palm_closed_front": "palm_close_front.jpg",
            "palm_closed_back": "palm_close_back.jpg"
        }
        
        for angle_key, img in images_dict.items():
            fname = view_filenames.get(angle_key, f"{angle_key}.jpg")
            img.save(target_dir / fname, "JPEG", quality=95)
            
        meta_file = target_dir / "metadata.json"
        meta_file.write_text(json.dumps(metadata_dict, indent=4))
        return True
    except Exception as e:
        print(f"[Local Save Error]: {e}")
        return False


def append_to_master_excel(user_id, gender, age, images_dict, final_dosha, probabilities, clinical_markers, is_doctor_mode=False):
    """
    Appends a new formatted row with embedded thumbnails to doctor_validation_prakriti_master.xlsx
    """
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.drawing.image import Image as OpenpyxlImage
        from pathlib import Path
        import io
        
        project_root = Path(__file__).resolve().parent.parent
        xlsx_paths = [
            project_root / "dataset" / "doctor_validation_prakriti_master.xlsx",
            project_root / "doctor_validation_prakriti_master.xlsx"
        ]
        
        for xlsx_path in xlsx_paths:
            if not xlsx_path.exists():
                continue
            wb = openpyxl.load_workbook(xlsx_path)
            ws = wb.active
            
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            col_map = {name: c for c, name in enumerate(headers, 1) if name}
            
            new_row = ws.max_row + 1
            ws.row_dimensions[new_row].height = 80
            
            # Demographics
            ws.cell(new_row, col_map["Participant ID"], value=f"{user_id}_{gender}")
            ws.cell(new_row, col_map["Participant Name"], value="Web App Submission")
            ws.cell(new_row, col_map["Age"], value=age)
            ws.cell(new_row, col_map["Sex"], value="F" if gender.lower() == "female" else "M")
            
            # Add thumbnails
            view_cols = {
                "palm_open_front": col_map.get("Palm Open Front (Palm)"),
                "palm_open_back": col_map.get("Palm Open Back (Skin/Veins)"),
                "palm_closed_front": col_map.get("Palm Closed Front (Fist)"),
                "palm_closed_back": col_map.get("Palm Closed Back (Outer Fist)")
            }
            
            for angle, col_idx in view_cols.items():
                if col_idx and angle in images_dict:
                    try:
                        img = images_dict[angle].copy()
                        img.thumbnail((90, 80))
                        buf = io.BytesIO()
                        img.convert("RGB").save(buf, format="JPEG", quality=85)
                        buf.seek(0)
                        ox_img = OpenpyxlImage(buf)
                        ox_img.width = img.width
                        ox_img.height = img.height
                        col_letter = openpyxl.utils.get_column_letter(col_idx)
                        ws.add_image(ox_img, f"{col_letter}{new_row}")
                    except Exception:
                        pass
            
            # Status and Predictions
            status = "Verified (Doctor-in-the-Loop Mode)" if is_doctor_mode else "Pending Doctor Validation"
            ws.cell(new_row, col_map["Validation Status"], value=status)
            ws.cell(new_row, col_map["ML Predicted Dosha"], value=final_dosha)
            
            k_pct = int(probabilities.get("Kapha", 0) * 100)
            p_pct = int(probabilities.get("Pitta", 0) * 100)
            v_pct = int(probabilities.get("Vata", 0) * 100)
            ws.cell(new_row, col_map["ML Confidence / Percentages"], value=f"k-{k_pct}%, p-{p_pct}%, v-{v_pct}%")
            
            # Clinical features
            if clinical_markers:
                ws.cell(new_row, col_map["Doctor Vein Visibility"], value=clinical_markers.get("vein_visibility"))
                ws.cell(new_row, col_map["Doctor Mole Presence"], value=clinical_markers.get("mole"))
                ws.cell(new_row, col_map["Doctor Skin Texture"], value=clinical_markers.get("skin_texture"))
                ws.cell(new_row, col_map["Doctor Skin Color"], value=clinical_markers.get("skin_color"))
                ws.cell(new_row, col_map["Doctor Palm Color"], value=clinical_markers.get("palm_color"))
                ws.cell(new_row, col_map["Doctor Hair Thickness"], value=clinical_markers.get("hair_thickness"))
                ws.cell(new_row, col_map["Doctor Hair Color"], value=clinical_markers.get("hair_color"))
                
            wb.save(xlsx_path)
        return True
    except Exception as e:
        print(f"[Excel Append Error]: {e}")
        return False


def upload_complete_analysis(user_id, gender, images_dict, probabilities, final_dosha, confidence, 
                             age=None, top_class=None, secondary_class=None, effective_markers=None, is_doctor_mode=False, *args, **kwargs):
    """
    Complete unified upload & save workflow across all 3 platforms:
      1. Supabase Storage (images) + Supabase Database (record + full metadata JSON)
      2. Local Folder (dataset/Wrist_Dataset/{id}/: 4 .jpg images + metadata.json)
      3. Master Excel (doctor_validation_prakriti_master.xlsx: appends row with thumbnails)
    """
    results = {
        "success": False,
        "record_id": None,
        "image_urls": {},
        "errors": []
    }
    
    # Extract clean clinical markers
    clean_markers = {}
    if effective_markers:
        for k, v in effective_markers.items():
            clean_markers[k.lower().replace(" ", "_")] = v[0] if isinstance(v, (tuple, list)) else v
            
    # Build complete metadata dict
    dual_dosha = f"{top_class}-{secondary_class}" if secondary_class and secondary_class != top_class else top_class
    metadata_dict = {
        "user_id": str(user_id),
        "gender": gender,
        "age": age,
        "ml_predicted_dosha": final_dosha,
        "dominant_dosha": top_class,
        "subdominant_dosha": secondary_class,
        "dual_dosha_label": dual_dosha,
        "probabilities": probabilities,
        "confidence": confidence,
        "clinical_features": clean_markers,
        "is_doctor_mode": is_doctor_mode,
        "validation_status": "Verified (Doctor-in-the-Loop Mode)" if is_doctor_mode else "Pending Doctor Validation",
        "source": "Streamlit Frontend Web App"
    }
    
    # 1. Upload images to Supabase storage if provided
    if images_dict:
        upload_result = upload_all_images(images_dict, user_id, gender)
        results["image_urls"] = upload_result.get("urls", {})
        if not upload_result["success"]:
            results["errors"].extend(upload_result.get("errors", []))
    
    # 2. Insert data into Supabase table
    insert_result = insert_wrist_data(user_id, gender, probabilities, dual_dosha or final_dosha, confidence, metadata_dict=metadata_dict)
    if not insert_result["success"]:
        results["errors"].append(f"Database insert failed: {insert_result['error']}")
    else:
        results["record_id"] = insert_result["record_id"]
    
    # 3. Save to Local Disk (dataset/Wrist_Dataset/)
    if images_dict:
        save_to_local_dataset(user_id, gender, images_dict, metadata_dict)
        
    # 4. Append to Master Excel
    append_to_master_excel(user_id, gender, age, images_dict, dual_dosha or final_dosha, probabilities, clean_markers, is_doctor_mode)
    
    results["success"] = True
    return results
